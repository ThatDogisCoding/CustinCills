"""Create Excel workbooks for browser and CSV test runs."""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import IllegalCharacterError


def _read_csv(path: Path) -> list[list[str]]:
    if not path.is_file():
        raise FileNotFoundError(f"CSV file does not exist: {path}")

    try:
        with path.open("r", encoding="utf-8-sig", newline="") as csv_file:
            rows = list(csv.reader(csv_file))
    except (OSError, UnicodeError, csv.Error) as exc:
        raise ValueError(f"Could not read CSV file {path.name}: {exc}") from exc

    if not rows or not any(any(value.strip() for value in row) for row in rows):
        raise ValueError(f"{path.name} is empty.")
    return rows


def _append_csv_rows(sheet, rows: list[list[str]]) -> None:
    try:
        for row_number, row in enumerate(rows, start=1):
            for column, value in enumerate(row, start=1):
                cell = sheet.cell(row=row_number, column=column, value=value)
                cell.data_type = "s"
    except IllegalCharacterError as exc:
        raise ValueError(
            "A CSV contains control characters that Excel cannot store."
        ) from exc


def _format_data_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    header_fill = PatternFill(fill_type="solid", fgColor="E31837")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top")

    for column_cells in sheet.columns:
        column_index = column_cells[0].column
        content_width = max(
            (len(str(cell.value or "")) for cell in column_cells),
            default=0,
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(content_width + 2, 12),
            50,
        )


def create_test_run_workbook(
    *,
    mindbody_csv: Path,
    axle_csv: Path,
    booked: int,
    shown: int,
    sold: int,
    output_path: Path,
) -> Path:
    """Write input CSV contents and reconciliation counts to an Excel workbook."""
    counts = (
        ("Booked", booked),
        ("Shown", shown),
        ("Sold", sold),
    )
    if any(type(value) is not int or value < 0 for _, value in counts):
        raise ValueError("Booked, Shown, and Sold counts must be nonnegative integers.")
    if output_path.resolve() in {mindbody_csv.resolve(), axle_csv.resolve()}:
        raise ValueError("The output workbook cannot overwrite an input CSV.")
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("The output file must use the .xlsx extension.")

    mindbody_rows = _read_csv(mindbody_csv)
    axle_rows = _read_csv(axle_csv)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Test Run Summary", "Value"])
    summary.append(["Generated", datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")])
    summary.append(["Mindbody CSV", mindbody_csv.name])
    summary.append(["Axle CSV", axle_csv.name])
    summary.append(["Booked", booked])
    summary.append(["Shown", shown])
    summary.append(["Sold", sold])
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 48
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="E31837")

    mindbody_sheet = workbook.create_sheet("Mindbody CSV")
    _append_csv_rows(mindbody_sheet, mindbody_rows)
    _format_data_sheet(mindbody_sheet)

    axle_sheet = workbook.create_sheet("Axle CSV")
    _append_csv_rows(axle_sheet, axle_rows)
    _format_data_sheet(axle_sheet)

    try:
        workbook.save(output_path)
    except OSError as exc:
        raise OSError(f"Could not save workbook to {output_path}: {exc}") from exc
    finally:
        workbook.close()

    return output_path
