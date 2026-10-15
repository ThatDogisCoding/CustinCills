import csv

import pytest
from openpyxl import load_workbook

from business_report.services.test_run_workbook import create_test_run_workbook


def _write_csv(path, rows):
    with path.open("w", encoding="utf-8", newline="") as csv_file:
        csv.writer(csv_file).writerows(rows)


def test_create_workbook_contains_counts_and_csv_sheets(tmp_path):
    mindbody_csv = tmp_path / "mindbody.csv"
    axle_csv = tmp_path / "axle.csv"
    output_path = tmp_path / "output" / "test-run.xlsx"
    _write_csv(mindbody_csv, [["Name", "Status"], ["Member A", "Active"]])
    _write_csv(axle_csv, [["Name", "Follow Up"], ["Member B", "Call"]])

    result = create_test_run_workbook(
        mindbody_csv=mindbody_csv,
        axle_csv=axle_csv,
        booked=10,
        shown=8,
        sold=3,
        output_path=output_path,
    )

    assert result == output_path
    workbook = load_workbook(output_path, data_only=True)
    try:
        assert workbook.sheetnames == ["Summary", "Mindbody CSV", "Axle CSV"]
        summary = workbook["Summary"]
        assert summary["A5"].value == "Booked"
        assert summary["B5"].value == 10
        assert summary["A6"].value == "Shown"
        assert summary["B6"].value == 8
        assert summary["A7"].value == "Sold"
        assert summary["B7"].value == 3
        assert workbook["Mindbody CSV"]["A2"].value == "Member A"
        assert workbook["Axle CSV"]["B2"].value == "Call"
    finally:
        workbook.close()


@pytest.mark.parametrize("count", [-1, 2.5, "2"])
def test_create_workbook_rejects_invalid_counts(tmp_path, count):
    mindbody_csv = tmp_path / "mindbody.csv"
    axle_csv = tmp_path / "axle.csv"
    _write_csv(mindbody_csv, [["Name"], ["Member A"]])
    _write_csv(axle_csv, [["Name"], ["Member B"]])

    with pytest.raises(ValueError, match="nonnegative integers"):
        create_test_run_workbook(
            mindbody_csv=mindbody_csv,
            axle_csv=axle_csv,
            booked=count,
            shown=1,
            sold=1,
            output_path=tmp_path / "test-run.xlsx",
        )


def test_create_workbook_rejects_empty_csv(tmp_path):
    mindbody_csv = tmp_path / "mindbody.csv"
    axle_csv = tmp_path / "axle.csv"
    mindbody_csv.write_text("", encoding="utf-8")
    _write_csv(axle_csv, [["Name"], ["Member B"]])

    with pytest.raises(ValueError, match="mindbody.csv is empty"):
        create_test_run_workbook(
            mindbody_csv=mindbody_csv,
            axle_csv=axle_csv,
            booked=1,
            shown=1,
            sold=0,
            output_path=tmp_path / "test-run.xlsx",
        )


def test_csv_formula_values_are_written_as_text(tmp_path):
    mindbody_csv = tmp_path / "mindbody.csv"
    axle_csv = tmp_path / "axle.csv"
    _write_csv(mindbody_csv, [["Name"], ["=1+1"]])
    _write_csv(axle_csv, [["Name"], ["Member B"]])
    output_path = tmp_path / "test-run.xlsx"

    create_test_run_workbook(
        mindbody_csv=mindbody_csv,
        axle_csv=axle_csv,
        booked=1,
        shown=1,
        sold=0,
        output_path=output_path,
    )

    workbook = load_workbook(output_path, data_only=False)
    try:
        cell = workbook["Mindbody CSV"]["A2"]
        assert cell.value == "=1+1"
        assert cell.data_type == "s"
    finally:
        workbook.close()


def test_workbook_cannot_overwrite_source_csv(tmp_path):
    mindbody_csv = tmp_path / "mindbody.csv"
    axle_csv = tmp_path / "axle.csv"
    _write_csv(mindbody_csv, [["Name"], ["Member A"]])
    _write_csv(axle_csv, [["Name"], ["Member B"]])

    with pytest.raises(ValueError, match="cannot overwrite an input CSV"):
        create_test_run_workbook(
            mindbody_csv=mindbody_csv,
            axle_csv=axle_csv,
            booked=1,
            shown=1,
            sold=0,
            output_path=mindbody_csv,
        )


def test_workbook_requires_xlsx_extension(tmp_path):
    mindbody_csv = tmp_path / "mindbody.csv"
    axle_csv = tmp_path / "axle.csv"
    _write_csv(mindbody_csv, [["Name"], ["Member A"]])
    _write_csv(axle_csv, [["Name"], ["Member B"]])

    with pytest.raises(ValueError, match="must use the .xlsx extension"):
        create_test_run_workbook(
            mindbody_csv=mindbody_csv,
            axle_csv=axle_csv,
            booked=1,
            shown=1,
            sold=0,
            output_path=tmp_path / "test-run.csv",
        )
